"""Alta de tareas desde un Excel: crea un ticket Task por fila y le carga horas."""

import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .carga import Carga, CargaError, formatear_duracion, parsear_duracion, registrar, segundos_por_dia, texto_a_adf
from .jira_client import JiraClient, JiraError

HOJA = "Tareas"
COLUMNAS = ["PROYECTO", "TITULO", "DESCRIPCION", "HORAS ESTIMADAS", "HORAS CARGADAS", "FECHA CARGA", "TICKET"]
ANCHOS = [12, 45, 60, 17, 17, 14, 14]
NOMBRES_TIPO_TAREA = ("task", "tarea")  # el tipo de issue a crear, según el idioma de Jira

# Campos de creación que este script completa o que Jira resuelve solo.
CAMPOS_CONOCIDOS = {"summary", "description", "issuetype", "project", "reporter", "assignee", "timetracking"}


@dataclass(frozen=True)
class Tarea:
    """Una fila del Excel."""

    fila: int
    proyecto: str
    titulo: str
    descripcion: str
    estimado_segundos: int | None
    cargado_segundos: int | None
    fecha_carga: date | None
    ticket: str  # vacío si hay que crearlo


@dataclass
class Plan:
    """Qué se va a hacer con una fila, después de validarla contra Jira."""

    tarea: Tarea
    crear: bool = False
    cargar: bool = False
    tipo_id: str = ""
    asignar: bool = False
    estimar_al_crear: bool = False
    error: str | None = None
    nota: str = ""

    def accion(self) -> str:
        if self.error:
            return f"ERROR: {self.error}"
        if self.crear and self.cargar:
            return "crear ticket + cargar horas"
        if self.crear:
            return "crear ticket"
        if self.cargar:
            return f"cargar horas en {self.tarea.ticket}"
        return f"nada que hacer ({self.nota})"


# --- Template ----------------------------------------------------------------


def crear_template(ruta: Path) -> None:
    if ruta.exists():
        raise CargaError(f"{ruta} ya existe; no lo piso. Borralo o elegí otro nombre.")
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA

    for col, (titulo, ancho) in enumerate(zip(COLUMNAS, ANCHOS), start=1):
        celda = ws.cell(1, col, titulo)
        celda.font = Font(name="Arial", bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="808080" if titulo == "TICKET" else "1F4E78")
        celda.alignment = Alignment(horizontal="center")
        ws.column_dimensions[celda.column_letter].width = ancho
    for fila in range(2, 501):
        for col in range(1, len(COLUMNAS) + 1):
            ws.cell(fila, col).font = Font(name="Arial")
        ws.cell(fila, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(fila, 6).number_format = "dd/mm/yyyy"

    validacion = DataValidation(
        type="date", operator="greaterThan", formula1="DATE(2020,1,1)",
        error="Poné una fecha válida", errorTitle="FECHA CARGA", showErrorMessage=True,
    )
    validacion.add("F2:F500")
    ws.add_data_validation(validacion)
    ws.freeze_panes = "A2"

    ayuda = wb.create_sheet("Instrucciones")
    lineas = [
        "Una fila de la hoja Tareas = un ticket Task nuevo, con horas cargadas en FECHA CARGA.",
        "",
        "PROYECTO: clave (AAA, BBB...) o ID del proyecto.",
        "TITULO: obligatorio. DESCRIPCION: opcional.",
        "HORAS ESTIMADAS / HORAS CARGADAS: 8, 2.5, 1h30m o 30m. Opcionales.",
        "FECHA CARGA: obligatoria si hay HORAS CARGADAS.",
        "TICKET: lo completa el script al crear el ticket. Si lo volvés a correr, esa fila no se crea de nuevo.",
        "  Si ponés a mano un ticket existente, solo se le cargan las horas.",
        "",
        "Ejemplo:",
    ]
    for n, texto in enumerate(lineas, start=1):
        ayuda.cell(n, 1, texto).font = Font(name="Arial")
    ejemplo = ["AAA", "Tarea de ejemplo", "Descripción de la tarea", 8, 4, "25/09/2026", ""]
    for col, (titulo, valor) in enumerate(zip(COLUMNAS, ejemplo), start=1):
        ayuda.cell(len(lineas) + 1, col, titulo).font = Font(name="Arial", bold=True)
        ayuda.cell(len(lineas) + 2, col, valor).font = Font(name="Arial")
    ayuda.column_dimensions["A"].width = 14

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


# --- Lectura -----------------------------------------------------------------


def leer_tareas(ruta: Path, zona: ZoneInfo) -> list[Tarea]:
    """Lee la hoja Tareas. Junta todos los errores de todas las filas antes de fallar."""
    if not ruta.exists():
        raise CargaError(f"No existe el archivo {ruta}")
    ws = _hoja(load_workbook(ruta, data_only=True))
    col = _columnas(ws)

    tareas, errores = [], []
    for fila in range(2, ws.max_row + 1):
        valores = {nombre: ws.cell(fila, col[nombre]).value for nombre in COLUMNAS}
        if all(_texto(v) == "" for v in valores.values()):
            continue
        try:
            tareas.append(_a_tarea(fila, valores, zona))
        except CargaError as e:
            errores.append(f"  fila {fila}: {e}")
    if errores:
        raise CargaError("El Excel tiene errores:\n" + "\n".join(errores))
    if not tareas:
        raise CargaError(f"La hoja {HOJA} no tiene ninguna fila con datos.")
    return tareas


def _a_tarea(fila: int, v: dict, zona: ZoneInfo) -> Tarea:
    tarea = Tarea(
        fila=fila,
        proyecto=_texto(v["PROYECTO"]).upper(),
        titulo=_texto(v["TITULO"]),
        descripcion=_texto(v["DESCRIPCION"]),
        estimado_segundos=_duracion(v["HORAS ESTIMADAS"], "HORAS ESTIMADAS"),
        cargado_segundos=_duracion(v["HORAS CARGADAS"], "HORAS CARGADAS"),
        fecha_carga=_fecha(v["FECHA CARGA"]),
        ticket=_texto(v["TICKET"]).upper(),
    )
    if not tarea.ticket and not tarea.proyecto:
        raise CargaError("falta PROYECTO.")
    if not tarea.ticket and not tarea.titulo:
        raise CargaError("falta TITULO.")
    if tarea.cargado_segundos and not tarea.fecha_carga:
        raise CargaError("hay HORAS CARGADAS pero falta FECHA CARGA.")
    if tarea.fecha_carga and not tarea.cargado_segundos:
        raise CargaError("hay FECHA CARGA pero faltan HORAS CARGADAS.")
    if tarea.fecha_carga and tarea.fecha_carga > datetime.now(zona).date():
        raise CargaError(f"FECHA CARGA {tarea.fecha_carga:%d/%m/%Y} es futura.")
    return tarea


def _hoja(wb: Workbook) -> Worksheet:
    return wb[HOJA] if HOJA in wb.sheetnames else wb.active


def _columnas(ws: Worksheet) -> dict[str, int]:
    """Nombre de columna -> número, tolerando mayúsculas, tildes y espacios de más."""
    encontradas = {_normalizar(c.value): c.column for c in ws[1] if c.value}
    faltan = [c for c in COLUMNAS if c not in encontradas]
    if faltan:
        raise CargaError(f"A la hoja {ws.title} le faltan las columnas: {', '.join(faltan)}")
    return {c: encontradas[c] for c in COLUMNAS}


def _normalizar(texto) -> str:
    sin_tildes = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(sin_tildes.upper().split())


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _duracion(valor, columna: str) -> int | None:
    """Número (horas), hora de Excel (4:30) o texto ('1h30m')."""
    if valor is None or _texto(valor) == "":
        return None
    try:
        if isinstance(valor, time):
            return valor.hour * 3600 + valor.minute * 60
        return parsear_duracion(str(valor))  # un número se toma como horas
    except CargaError as e:
        raise CargaError(f"{columna}: {e}") from None


def _fecha(valor) -> date | None:
    if valor is None or _texto(valor) == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise CargaError(f"FECHA CARGA inválida: {texto!r}. Usá una fecha de Excel o DD/MM/AAAA.")


# --- Contra Jira -------------------------------------------------------------


@dataclass
class _InfoProyecto:
    tipo_id: str = ""
    permite_estimacion: bool = False
    permite_asignar: bool = False
    error: str | None = None


def planificar(cliente: JiraClient, tareas: list[Tarea], zona: ZoneInfo) -> tuple[list[Plan], str]:
    """Decide qué hacer con cada fila y lo valida contra Jira. Devuelve los planes y tu accountId."""
    yo = cliente.usuario_actual()["accountId"]
    proyectos = {t.proyecto: _info_proyecto(cliente, t.proyecto) for t in tareas if not t.ticket}

    planes = []
    for t in tareas:
        plan = Plan(t, cargar=bool(t.cargado_segundos))
        if not t.ticket:
            info = proyectos[t.proyecto]
            plan.crear, plan.tipo_id, plan.asignar = True, info.tipo_id, info.permite_asignar
            # Si la pantalla de creación no tiene la estimación, se setea editando el ticket recién creado.
            plan.estimar_al_crear = info.permite_estimacion
            plan.error = info.error
        elif plan.cargar:
            # El ticket ya existe (lo creó una corrida anterior o lo pusiste a mano):
            # cargar solo si todavía no tenés horas ahí ese día.
            try:
                ya = segundos_por_dia(cliente, t.ticket, t.fecha_carga, zona, yo).get(t.fecha_carga, 0)
            except JiraError as e:
                plan.error = str(e)
            else:
                if ya:
                    plan.cargar = False
                    plan.nota = f"ya tiene {formatear_duracion(ya)} tuyas el {t.fecha_carga:%d/%m}"
        else:
            plan.nota = f"{t.ticket} ya creado y sin horas para cargar"
        planes.append(plan)
    return planes, yo


def _info_proyecto(cliente: JiraClient, proyecto: str) -> _InfoProyecto:
    try:
        tipos = cliente.tipos_de_issue(proyecto)
        tipo = next((t for t in tipos if t.get("name", "").strip().lower() in NOMBRES_TIPO_TAREA), None)
        if not tipo:
            nombres = ", ".join(t.get("name", "?") for t in tipos) or "ninguno"
            return _InfoProyecto(error=f"el proyecto {proyecto} no tiene tipo Task (tiene: {nombres})")
        campos = cliente.campos_de_creacion(proyecto, tipo["id"])
    except JiraError as e:
        return _InfoProyecto(error=str(e))

    ids = {c.get("fieldId") or c.get("key") for c in campos}
    obligatorios = [
        c.get("name") or c.get("fieldId")
        for c in campos
        if c.get("required") and not c.get("hasDefaultValue") and (c.get("fieldId") or c.get("key")) not in CAMPOS_CONOCIDOS
    ]
    if obligatorios:
        return _InfoProyecto(error=f"el proyecto {proyecto} pide campos obligatorios no soportados: {', '.join(obligatorios)}")
    return _InfoProyecto(
        tipo_id=str(tipo["id"]),
        permite_estimacion="timetracking" in ids,
        permite_asignar="assignee" in ids,
    )


def ejecutar(
    cliente: JiraClient,
    planes: list[Plan],
    ruta_excel: Path,
    yo: str,
    zona: ZoneInfo,
    hora: time,
    ajustar_estimacion: bool = True,
) -> int:
    """Crea los tickets y carga las horas. Escribe cada ticket nuevo en el Excel apenas se crea.

    Devuelve la cantidad de filas que fallaron.
    """
    wb = load_workbook(ruta_excel)
    ws = _hoja(wb)
    col_ticket = _columnas(ws)["TICKET"]
    fallidas = 0

    for plan in planes:
        if plan.error or not (plan.crear or plan.cargar):
            continue
        t = plan.tarea
        ticket = t.ticket
        try:
            if plan.crear:
                ticket = _crear_ticket(cliente, plan, yo)
                celda = ws.cell(t.fila, col_ticket, ticket)
                celda.hyperlink = f"{cliente.url}/browse/{ticket}"
                celda.font = Font(name="Arial", color="0563C1", underline="single")
                wb.save(ruta_excel)  # guardar ya: si algo falla después, no se vuelve a crear
                print(f"  fila {t.fila}: creado {ticket} «{t.titulo[:50]}»")
                if t.estimado_segundos and not plan.estimar_al_crear:
                    _estimar_despues(cliente, ticket, t)
            if plan.cargar:
                carga = Carga(
                    ticket=ticket,
                    segundos=t.cargado_segundos,
                    inicio=datetime.combine(t.fecha_carga, hora, tzinfo=zona),
                    comentario="",
                    origen=f"fila {t.fila}",
                )
                registrar(cliente, carga, ajustar_estimacion)
                print(f"  fila {t.fila}: cargadas {formatear_duracion(t.cargado_segundos)} en {ticket} el {t.fecha_carga:%d/%m}")
        except JiraError as e:
            fallidas += 1
            print(f"  fila {t.fila}: ERROR {e}")
    return fallidas


def _crear_ticket(cliente: JiraClient, plan: Plan, yo: str) -> str:
    t = plan.tarea
    proyecto = {"id": t.proyecto} if t.proyecto.isdigit() else {"key": t.proyecto}
    campos: dict = {"project": proyecto, "issuetype": {"id": plan.tipo_id}, "summary": t.titulo}
    if t.descripcion:
        campos["description"] = texto_a_adf(t.descripcion)
    if t.estimado_segundos and plan.estimar_al_crear:
        campos["timetracking"] = {"originalEstimate": formatear_duracion(t.estimado_segundos)}
    if plan.asignar:
        campos["assignee"] = {"accountId": yo}
    return cliente.crear_issue(campos)["key"]


def _estimar_despues(cliente: JiraClient, ticket: str, t: Tarea) -> None:
    """Setea la estimación original editando el ticket. Si Jira no deja, avisa y sigue."""
    estimado = formatear_duracion(t.estimado_segundos)
    try:
        cliente.editar_issue(ticket, {"timetracking": {"originalEstimate": estimado}})
        print(f"  fila {t.fila}: estimación {estimado} en {ticket}")
    except JiraError as e:
        print(f"  fila {t.fila}: AVISO no se pudo poner la estimación {estimado} en {ticket}: {e}")
