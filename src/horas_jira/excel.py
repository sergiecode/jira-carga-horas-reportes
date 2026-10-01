"""Generación del Excel de horas (hojas Detalle y Resumen)."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .periodo import Periodo
from .registros import Registro

FUENTE = "Arial"
DIAS_SEMANA = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]

# Formato de horas: dos decimales y los ceros como "-".
FORMATO_HORAS = '0.00;-0.00;"-"'

RELLENO_ENCABEZADO = PatternFill("solid", fgColor="1F4E78")
RELLENO_EDITABLE = PatternFill("solid", fgColor="FFFF00")
RELLENO_FIN_DE_SEMANA = PatternFill("solid", fgColor="D9D9D9")
RELLENO_FALTAN_HORAS = PatternFill("solid", fgColor="FFC7CE", bgColor="FFC7CE")
BORDE_SUPERIOR = Border(top=Side(style="thin"))

# Columnas de la hoja Detalle. Las fórmulas del Resumen dependen de estas letras.
COLUMNAS_DETALLE = ["Fecha", "Persona", "Ticket", "Resumen", "Proyecto", "Horas", "Comentario"]
ANCHOS_DETALLE = [12, 26, 14, 50, 26, 9, 60]
COL_FECHA, COL_PERSONA, COL_HORAS = "A", "B", "F"

# Ubicación de las cosas en la hoja Resumen.
CELDA_MINIMO = "B2"
CELDA_MINIMO_ABS = "$B$2"
FILA_ENCABEZADO_RESUMEN = 5


def nombre_archivo(periodo: Periodo) -> str:
    return f"horas_{periodo.desde:%Y-%m-%d}_a_{periodo.hasta:%Y-%m-%d}.xlsx"


def generar_excel(
    registros: list[Registro],
    periodo: Periodo,
    *,
    carpeta_salida: Path,
    jira_url: str,
    horas_minimas_dia: float,
    equipo: list[str],
    hoy: date | None = None,
) -> Path:
    """Arma el Excel y lo guarda en `carpeta_salida`. Devuelve la ruta del archivo."""
    wb = Workbook()
    detalle = wb.active
    detalle.title = "Detalle"
    _hoja_detalle(detalle, registros, jira_url)

    personas = sorted({r.persona for r in registros} | set(equipo), key=str.casefold)
    _hoja_resumen(wb.create_sheet("Resumen"), personas, periodo, horas_minimas_dia, hoy)

    wb.calculation.fullCalcOnLoad = True
    carpeta_salida.mkdir(parents=True, exist_ok=True)
    ruta = carpeta_salida / nombre_archivo(periodo)
    wb.save(ruta)
    return ruta


def _hoja_detalle(ws: Worksheet, registros: list[Registro], jira_url: str) -> None:
    _encabezado(ws, 1, COLUMNAS_DETALLE)
    for i, ancho in enumerate(ANCHOS_DETALLE, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    fuente = Font(name=FUENTE)
    fuente_link = Font(name=FUENTE, color="0563C1", underline="single")
    for fila, r in enumerate(registros, start=2):
        valores = [r.fecha, r.persona, r.ticket, r.resumen, r.proyecto, r.horas, r.comentario]
        for col, valor in enumerate(valores, start=1):
            ws.cell(fila, col, valor).font = fuente
        ws.cell(fila, 1).number_format = "dd/mm/yyyy"
        ws.cell(fila, 6).number_format = "0.00"
        ticket = ws.cell(fila, 3)
        ticket.hyperlink = f"{jira_url}/browse/{r.ticket}"
        ticket.font = fuente_link
        ws.cell(fila, 7).alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS_DETALLE))}{max(len(registros) + 1, 2)}"


def _hoja_resumen(
    ws: Worksheet,
    personas: list[str],
    periodo: Periodo,
    horas_minimas_dia: float,
    hoy: date | None,
) -> None:
    dias = periodo.dias()
    fuente = Font(name=FUENTE)

    # Título, mínimo editable y leyenda.
    ws["A1"] = f"Horas cargadas del {periodo.desde:%d/%m/%Y} al {periodo.hasta:%d/%m/%Y}"
    ws["A1"].font = Font(name=FUENTE, bold=True, size=14)
    ws["A2"] = "Mínimo de horas por día:"
    ws["A2"].font = Font(name=FUENTE, bold=True)
    ws[CELDA_MINIMO] = horas_minimas_dia
    ws[CELDA_MINIMO].font = Font(name=FUENTE, bold=True)
    ws[CELDA_MINIMO].fill = RELLENO_EDITABLE
    ws["C2"] = "← editable"
    ws["C2"].font = Font(name=FUENTE, italic=True, color="808080")
    ws["A3"] = (
        "En rojo: días hábiles con menos horas que el mínimo (el día de hoy no se marca). "
        "En gris: fines de semana."
    )
    ws["A3"].font = Font(name=FUENTE, italic=True, size=9, color="808080")

    # Encabezado: Persona | lun 28/09 | ... | Total
    fila_enc = FILA_ENCABEZADO_RESUMEN
    titulos = ["Persona"] + [f"{DIAS_SEMANA[d.weekday()]} {d:%d/%m}" for d in dias] + ["Total"]
    _encabezado(ws, fila_enc, titulos)
    col_total = len(dias) + 2
    letra_total = get_column_letter(col_total)

    # Una fila por persona, con SUMIFS contra la hoja Detalle.
    primera = fila_enc + 1
    ultima = fila_enc + len(personas)
    for fila, persona in enumerate(personas, start=primera):
        ws.cell(fila, 1, persona).font = fuente
        for col, dia in enumerate(dias, start=2):
            celda = ws.cell(fila, col, _formula_horas_dia(f"$A{fila}", dia))
            celda.font = fuente
            celda.number_format = FORMATO_HORAS
        total = ws.cell(fila, col_total, f"=SUM(B{fila}:{get_column_letter(col_total - 1)}{fila})")
        total.font = Font(name=FUENTE, bold=True)
        total.number_format = FORMATO_HORAS

    # Fila "Total equipo".
    fila_tot = ultima + 1
    ws.cell(fila_tot, 1, "Total equipo")
    for col in range(1, col_total + 1):
        celda = ws.cell(fila_tot, col)
        if col > 1:
            letra = get_column_letter(col)
            celda.value = f"=SUM({letra}{primera}:{letra}{ultima})" if personas else 0
            celda.number_format = FORMATO_HORAS
        celda.font = Font(name=FUENTE, bold=True)
        celda.border = BORDE_SUPERIOR

    # Colores por día: gris fijo los fines de semana, rojo condicional los hábiles (salvo hoy).
    for col, dia in enumerate(dias, start=2):
        letra = get_column_letter(col)
        if dia.weekday() >= 5:
            for fila in range(primera, fila_tot + 1):
                ws.cell(fila, col).fill = RELLENO_FIN_DE_SEMANA
        elif dia != hoy and personas:
            ws.conditional_formatting.add(
                f"{letra}{primera}:{letra}{ultima}",
                FormulaRule(
                    formula=[f"{letra}{primera}<{CELDA_MINIMO_ABS}"],
                    fill=RELLENO_FALTAN_HORAS,
                    font=Font(color="9C0006"),
                ),
            )

    ws.column_dimensions["A"].width = 28
    for col in range(2, col_total):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.column_dimensions[letra_total].width = 10
    ws.freeze_panes = ws.cell(primera, 2)


def _formula_horas_dia(ref_persona: str, dia: date) -> str:
    """Suma de horas de Detalle para una persona y un día."""
    return (
        f"=SUMIFS(Detalle!${COL_HORAS}:${COL_HORAS},"
        f"Detalle!${COL_PERSONA}:${COL_PERSONA},{ref_persona},"
        f"Detalle!${COL_FECHA}:${COL_FECHA},DATE({dia.year},{dia.month},{dia.day}))"
    )


def _encabezado(ws: Worksheet, fila: int, titulos: list[str]) -> None:
    fuente = Font(name=FUENTE, bold=True, color="FFFFFF")
    for col, titulo in enumerate(titulos, start=1):
        celda = ws.cell(fila, col, titulo)
        celda.font = fuente
        celda.fill = RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center")
