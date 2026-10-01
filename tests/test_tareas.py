"""Tests de crear-tareas, con un Jira falso que guarda estado (tickets y worklogs creados)."""

import json
from datetime import date

import pytest
import requests
from openpyxl import load_workbook

from horas_jira import cli_tareas
from horas_jira.carga import CargaError
from horas_jira.tareas import COLUMNAS, crear_template, leer_tareas

URL = "https://ejemplo.atlassian.net"


def respuesta(datos: dict | None, status: int = 200) -> requests.Response:
    r = requests.Response()
    r.status_code = status
    r._content = json.dumps(datos).encode() if datos is not None else b""
    return r


class JiraFalso:
    """AAA: la pantalla de creación no tiene timetracking (hay que editarlo después). BBB: sí lo tiene."""

    def __init__(self):
        self.issues: dict[str, dict] = {}
        self.ediciones: list[tuple[str, dict]] = []
        self.worklogs: dict[str, list[dict]] = {}
        self.fallar_worklogs = False

    def __call__(self, metodo, url, **kwargs):
        ruta = url.removeprefix(URL)
        partes = ruta.split("/")
        if ruta == "/rest/api/3/myself":
            return respuesta({"accountId": "yo"})
        if ruta.startswith("/rest/api/3/issue/createmeta/"):
            # /rest/api/3/issue/createmeta/{proyecto}/issuetypes[/{tipo}]
            proyecto = partes[6]
            if proyecto == "NOPE":
                return respuesta({}, status=404)
            if len(partes) == 8:
                return respuesta({"issueTypes": [{"id": "901", "name": "Otra Task"}, {"id": "902", "name": "Task"}]})
            campos = [{"fieldId": f, "required": f in ("summary", "project", "issuetype")}
                      for f in ("summary", "project", "issuetype", "description", "assignee")]
            if proyecto == "BBB":
                campos.append({"fieldId": "timetracking", "required": False})
            return respuesta({"fields": campos})
        if ruta == "/rest/api/3/issue" and metodo == "POST":
            campos = kwargs["json"]["fields"]
            proyecto = campos["project"]["key"]
            key = f"{proyecto}-{100 + len(self.issues)}"
            self.issues[key] = campos
            return respuesta({"id": "1", "key": key}, status=201)
        if metodo == "PUT":
            self.ediciones.append((partes[5], kwargs["json"]["fields"]))
            return respuesta(None, status=204)
        if ruta.endswith("/worklog"):
            key = partes[5]
            if metodo == "GET":
                todos = self.worklogs.get(key, [])
                return respuesta({"total": len(todos), "worklogs": todos})
            if self.fallar_worklogs:
                return respuesta({"errorMessages": ["algo salió mal"]}, status=400)
            wl = {**kwargs["json"], "author": {"accountId": "yo"}, "params": kwargs["params"]}
            self.worklogs.setdefault(key, []).append(wl)
            return respuesta({"id": str(len(self.worklogs[key]))}, status=201)
        raise AssertionError(f"ruta inesperada: {metodo} {ruta}")


@pytest.fixture
def jira(monkeypatch):
    falso = JiraFalso()
    monkeypatch.setattr(requests.Session, "request", falso)
    monkeypatch.setattr("horas_jira.config.load_dotenv", lambda: None)
    for var, valor in {"JIRA_URL": URL, "JIRA_EMAIL": "yo@ejemplo.com", "JIRA_API_TOKEN": "t"}.items():
        monkeypatch.setenv(var, valor)
    return falso


def excel(tmp_path, filas: list[list]):
    ruta = tmp_path / "tareas.xlsx"
    crear_template(ruta)
    wb = load_workbook(ruta)
    for n, fila in enumerate(filas, start=2):
        for col, valor in enumerate(fila, start=1):
            wb["Tareas"].cell(n, col, valor)
    wb.save(ruta)
    return ruta


def correr(monkeypatch, *argv):
    monkeypatch.setattr("sys.argv", ["crear-tareas", *argv])
    return cli_tareas.main()


def tickets_en_excel(ruta):
    ws = load_workbook(ruta)["Tareas"]
    valores = [ws.cell(f, COLUMNAS.index("TICKET") + 1).value for f in range(2, ws.max_row + 1)]
    return [v for v in valores if v]


# --- Template y lectura --------------------------------------------------------


def test_template(tmp_path):
    ruta = tmp_path / "tareas.xlsx"
    crear_template(ruta)
    wb = load_workbook(ruta)
    assert [c.value for c in wb["Tareas"][1]] == COLUMNAS
    assert "Instrucciones" in wb.sheetnames
    with pytest.raises(CargaError, match="ya existe"):
        crear_template(ruta)
    with pytest.raises(CargaError, match="ninguna fila"):
        leer_tareas(ruta, __import__("zoneinfo").ZoneInfo("UTC"))


def test_lectura_acepta_varios_formatos(tmp_path):
    from zoneinfo import ZoneInfo

    ruta = excel(tmp_path, [
        ["aaa", "A", "desc", 8, "1h30m", date(2026, 9, 25)],
        ["AAA", "B", None, "2,5", None, None],
        ["AAA", "C", None, None, 1, "28/09/2026"],
    ])
    a, b, c = leer_tareas(ruta, ZoneInfo("America/Argentina/Buenos_Aires"))
    assert (a.proyecto, a.estimado_segundos, a.cargado_segundos, a.fecha_carga) == ("AAA", 28800, 5400, date(2026, 9, 25))
    assert (b.estimado_segundos, b.cargado_segundos, b.fecha_carga) == (9000, None, None)
    assert c.fecha_carga == date(2026, 9, 28)


def test_lectura_junta_todos_los_errores(tmp_path):
    from zoneinfo import ZoneInfo

    ruta = excel(tmp_path, [
        ["AAA", "", None, None, None, None],                # falta título
        ["AAA", "B", None, None, "4h", None],               # carga sin fecha
        ["AAA", "C", None, "mucho", None, None],            # duración inválida
        ["AAA", "D", None, None, "1h", date(2999, 1, 1)],   # fecha futura
        ["AAA", "E", None, None, None, None],               # OK
    ])
    with pytest.raises(CargaError) as e:
        leer_tareas(ruta, ZoneInfo("UTC"))
    for fila in (2, 3, 4, 5):
        assert f"fila {fila}" in str(e.value)
    assert "fila 6" not in str(e.value)


# --- Flujo completo ------------------------------------------------------------


def test_crea_task_estima_despues_y_carga_horas(jira, monkeypatch, tmp_path):
    ruta = excel(tmp_path, [["AAA", "Mi tarea", "Línea 1\nLínea 2", 8, 4, date(2026, 9, 25)]])
    assert correr(monkeypatch, str(ruta), "--si") == 0

    (key, campos), = jira.issues.items()
    assert key == "AAA-100"
    assert campos["issuetype"] == {"id": "902"}  # Task, no Otra Task
    assert campos["assignee"] == {"accountId": "yo"}
    assert len(campos["description"]["content"]) == 2
    assert "timetracking" not in campos  # AAA no lo tiene en la pantalla de creación...
    assert jira.ediciones == [("AAA-100", {"timetracking": {"originalEstimate": "8h"}})]  # ...se edita después

    (wl,) = jira.worklogs["AAA-100"]
    assert wl["started"] == "2026-09-25T09:00:00.000-0300"
    assert wl["timeSpentSeconds"] == 4 * 3600
    assert tickets_en_excel(ruta) == ["AAA-100"]


def test_si_la_pantalla_lo_permite_estima_al_crear(jira, monkeypatch, tmp_path):
    ruta = excel(tmp_path, [["BBB", "Otra", None, "1h30m", None, None]])
    assert correr(monkeypatch, str(ruta), "--si") == 0
    assert jira.issues["BBB-100"]["timetracking"] == {"originalEstimate": "1h 30m"}
    assert jira.ediciones == []


def test_volver_a_correr_no_duplica(jira, monkeypatch, tmp_path):
    ruta = excel(tmp_path, [["AAA", "Mi tarea", None, None, 4, date(2026, 9, 25)]])
    assert correr(monkeypatch, str(ruta), "--si") == 0
    assert correr(monkeypatch, str(ruta), "--si") == 0
    assert len(jira.issues) == 1
    assert len(jira.worklogs["AAA-100"]) == 1


def test_si_falla_la_carga_la_reintenta_sin_recrear(jira, monkeypatch, tmp_path):
    ruta = excel(tmp_path, [["AAA", "Mi tarea", None, None, 4, date(2026, 9, 25)]])
    jira.fallar_worklogs = True
    assert correr(monkeypatch, str(ruta), "--si") == 1
    assert tickets_en_excel(ruta) == ["AAA-100"]  # el ticket quedó anotado igual

    jira.fallar_worklogs = False
    assert correr(monkeypatch, str(ruta), "--si") == 0
    assert len(jira.issues) == 1
    assert len(jira.worklogs["AAA-100"]) == 1


def test_ticket_existente_solo_carga_horas(jira, monkeypatch, tmp_path):
    ruta = excel(tmp_path, [[None, None, None, None, "2h", date(2026, 9, 25), "AAA-999"]])
    assert correr(monkeypatch, str(ruta), "--si", "--mantener-estimacion") == 0
    assert jira.issues == {}
    assert jira.worklogs["AAA-999"][0]["params"] == {"adjustEstimate": "leave"}


def test_dry_run_y_proyecto_inexistente(jira, monkeypatch, tmp_path, capsys):
    ruta = excel(tmp_path, [["AAA", "Bien", None, 1, 1, date(2026, 9, 25)], ["NOPE", "Mal", None, None, None, None]])
    assert correr(monkeypatch, str(ruta), "--dry-run") == 1
    salida = capsys.readouterr().out
    assert "crear ticket + cargar horas" in salida
    assert "ERROR: No existe" in salida
    assert jira.issues == {} and jira.worklogs == {}
