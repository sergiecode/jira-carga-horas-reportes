"""Tests de punta a punta con la API de Jira mockeada.

Se reemplaza requests.Session.request por un servidor falso que simula:
paginación de /search/jql (nextPageToken) y de /worklog (startAt/total), un 429 con
Retry-After, worklogs fuera del período y un worklog que cambia de día al pasar a la zona local.
"""

import json
from datetime import date
from zoneinfo import ZoneInfo

import pytest
import requests
from openpyxl import load_workbook

from horas_jira.excel import generar_excel
from horas_jira.jira_client import JiraClient, JiraError
from horas_jira.periodo import Periodo, ultimos_dias
from horas_jira.registros import adf_a_texto, armar_jql, obtener_registros

ZONA = ZoneInfo("America/Argentina/Buenos_Aires")
HOY = date(2026, 10, 1)  # jueves
PERIODO = ultimos_dias(7, ZONA, HOY)  # vie 25/09 a jue 01/10
URL = "https://ejemplo.atlassian.net"


def adf(texto: str) -> dict:
    return {"type": "doc", "version": 1, "content": [{"type": "paragraph", "content": [{"type": "text", "text": texto}]}]}


def wl(persona: str, started: str, horas: float, comentario: str | None = None) -> dict:
    return {
        "author": {"displayName": persona},
        "started": started,
        "timeSpentSeconds": int(horas * 3600),
        "comment": adf(comentario) if comentario else None,
    }


def issue(key: str, resumen: str) -> dict:
    return {"key": key, "fields": {"summary": resumen, "project": {"name": "Proyecto Demo"}}}


WORKLOGS = {
    "DEMO-1": [
        wl("Ana", "2026-09-28T09:00:00.000-0300", 8, "Desarrollo del login"),
        wl("Ana", "2026-09-29T09:00:00.000-0300", 4),
        wl("Beto", "2026-09-28T14:00:00.000-0300", 3),
        wl("Ana", "2026-09-24T10:00:00.000-0300", 6),  # antes del período
    ],
    "DEMO-2": [
        wl("Beto", "2026-09-28T09:00:00.000-0300", 5),
        wl("Ana", "2026-09-29T14:00:00.000-0300", 4),
        wl("Beto", "2026-10-02T09:00:00.000-0300", 8),  # después del período
    ],
    # 01:30 UTC del 01/10 = 22:30 del 30/09 en Buenos Aires.
    "DEMO-3": [wl("Ana", "2026-10-01T01:30:00.000+0000", 8)],
}


class JiraFalso:
    def __init__(self):
        self.llamadas: list[tuple[str, str]] = []
        self.ya_devolvio_429 = False

    def __call__(self, metodo, url, **kwargs):
        ruta = url.removeprefix(URL)
        self.llamadas.append((metodo, ruta))

        if ruta == "/rest/api/3/search/jql":
            if kwargs["json"].get("nextPageToken") is None:
                return respuesta({"issues": [issue("DEMO-1", "Login"), issue("DEMO-2", "Reportes")], "nextPageToken": "p2"})
            return respuesta({"issues": [issue("DEMO-3", "Deploy")], "isLast": True})

        key = ruta.split("/")[5]
        if key == "DEMO-3" and not self.ya_devolvio_429:
            self.ya_devolvio_429 = True
            return respuesta({}, status=429, headers={"Retry-After": "0"})

        # Páginas de 2 worklogs para forzar la paginación con startAt/total.
        todos = WORKLOGS[key]
        start = kwargs["params"]["startAt"]
        return respuesta({"startAt": start, "total": len(todos), "worklogs": todos[start : start + 2]})


def respuesta(datos: dict, status: int = 200, headers: dict | None = None) -> requests.Response:
    r = requests.Response()
    r.status_code = status
    r._content = json.dumps(datos).encode()
    r.headers.update(headers or {})
    return r


@pytest.fixture
def jira(monkeypatch):
    falso = JiraFalso()
    # Una instancia callable no se "bindea" como método: recibe (metodo, url, ...) sin la sesión.
    monkeypatch.setattr(requests.Session, "request", falso)
    return falso


@pytest.fixture
def registros(jira):
    cliente = JiraClient(URL, "yo@ejemplo.com", "token")
    return obtener_registros(cliente, ["AAA", "BBB"], PERIODO, ZONA)


# --- Registros ---------------------------------------------------------------


def test_filtra_por_periodo_y_zona_horaria(registros):
    por_dia = {(r.persona, r.fecha): 0.0 for r in registros}
    for r in registros:
        por_dia[(r.persona, r.fecha)] += r.horas
    assert por_dia == {
        ("Ana", date(2026, 9, 28)): 8,
        ("Ana", date(2026, 9, 29)): 8,
        ("Ana", date(2026, 9, 30)): 8,  # el de DEMO-3, convertido a hora local
        ("Beto", date(2026, 9, 28)): 8,
    }
    assert len(registros) == 6


def test_pagina_y_reintenta(jira, registros):
    assert jira.llamadas.count(("POST", "/rest/api/3/search/jql")) == 2
    assert jira.llamadas.count(("GET", "/rest/api/3/issue/DEMO-1/worklog")) == 2  # 4 worklogs
    assert jira.llamadas.count(("GET", "/rest/api/3/issue/DEMO-3/worklog")) == 2  # 429 + reintento


def test_comentario_adf(registros):
    login = next(r for r in registros if r.comentario)
    assert login.comentario == "Desarrollo del login"
    assert login.ticket == "DEMO-1"
    assert login.resumen == "Login"


def test_jql_con_y_sin_proyectos():
    fechas = 'worklogDate >= "2026-09-24" AND worklogDate <= "2026-10-02"'
    assert armar_jql(["AAA", "10001"], PERIODO) == f"project in (AAA, 10001) AND {fechas}"
    assert armar_jql([], PERIODO) == fechas  # todos los proyectos visibles


def test_adf_a_texto():
    doc = {
        "type": "doc",
        "content": [
            {"type": "paragraph", "content": [{"type": "text", "text": "Hola "}, {"type": "mention", "attrs": {"text": "@Ana"}}]},
            {"type": "bulletList", "content": [{"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "item"}]}]}]},
        ],
    }
    assert adf_a_texto(doc) == "Hola @Ana\nitem"
    assert adf_a_texto(None) == ""


def test_401_mensaje_claro(monkeypatch):
    monkeypatch.setattr(requests.Session, "request", lambda *a, **k: respuesta({}, status=401))
    with pytest.raises(JiraError, match="credenciales"):
        JiraClient(URL, "yo@ejemplo.com", "malo").obtener_worklogs("DEMO-1", 0)


def test_403_mensaje_claro(monkeypatch):
    monkeypatch.setattr(requests.Session, "request", lambda *a, **k: respuesta({}, status=403))
    with pytest.raises(JiraError, match="permisos"):
        list(JiraClient(URL, "yo@ejemplo.com", "token").buscar_issues("x", []))


# --- Excel -------------------------------------------------------------------


@pytest.fixture
def libro(registros, tmp_path):
    ruta = generar_excel(
        registros,
        PERIODO,
        carpeta_salida=tmp_path,
        jira_url=URL,
        horas_minimas_dia=8,
        equipo=["Ana", "Beto", "Carla"],
        hoy=HOY,
    )
    assert ruta.name == "horas_2026-09-25_a_2026-10-01.xlsx"
    return load_workbook(ruta)


def test_hoja_detalle(libro):
    ws = libro["Detalle"]
    assert [c.value for c in ws[1]] == ["Fecha", "Persona", "Ticket", "Resumen", "Proyecto", "Horas", "Comentario"]
    assert ws.max_row == 7
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:G7"
    assert ws["A2"].is_date
    assert ws["C2"].hyperlink.target.startswith(f"{URL}/browse/DEMO-")
    assert ws["B2"].font.name == "Arial"


def test_hoja_resumen(libro):
    ws = libro["Resumen"]
    assert ws["A1"].value == "Horas cargadas del 25/09/2026 al 01/10/2026"
    assert ws["B2"].value == 8
    assert ws["B2"].fill.fgColor.rgb == "00FFFF00"

    assert [c.value for c in ws[5]] == [
        "Persona", "vie 25/09", "sáb 26/09", "dom 27/09", "lun 28/09", "mar 29/09", "mié 30/09", "jue 01/10", "Total",
    ]
    assert [ws.cell(f, 1).value for f in range(6, 10)] == ["Ana", "Beto", "Carla", "Total equipo"]

    assert ws["E6"].value == "=SUMIFS(Detalle!$F:$F,Detalle!$B:$B,$A6,Detalle!$A:$A,DATE(2026,9,28))"
    assert ws["I6"].value == "=SUM(B6:H6)"
    assert ws["B9"].value == "=SUM(B6:B8)"
    assert ws["E6"].number_format == '0.00;-0.00;"-"'

    # Rojo condicional solo en días hábiles que no son hoy: vie, lun, mar, mié.
    rangos = sorted(str(cf.sqref) for cf in ws.conditional_formatting)
    assert rangos == ["B6:B8", "E6:E8", "F6:F8", "G6:G8"]
    # Fines de semana en gris.
    assert ws["C7"].fill.fgColor.rgb == "00D9D9D9"
    assert ws["D9"].fill.fgColor.rgb == "00D9D9D9"


def test_sin_registros_igual_genera(tmp_path):
    ruta = generar_excel(
        [], Periodo(HOY, HOY), carpeta_salida=tmp_path, jira_url=URL, horas_minimas_dia=8, equipo=[], hoy=HOY
    )
    assert load_workbook(ruta)["Resumen"]["A6"].value == "Total equipo"
